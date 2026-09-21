import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel(filepath):
    wb = openpyxl.Workbook()
    
    # Styles
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    sub_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    bold_font = Font(name="Arial", size=10, bold=True)
    regular_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # -------------------------------------------------------------
    # 1. Readme_Portada
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Readme_Portada"
    ws1.views.sheetView[0].showGridLines = True
    
    data1 = [
        ["AutoITBA S.A. - Base de Datos de Parámetros", ""],
        ["Resumen Ejecutivo", ""],
        ["Terminal automotriz", "AutoITBA S.A. (radicada en Zárate)"],
        ["Horizonte de planificación", "5 años (Año 1 a Año 5)"],
        ["", ""],
        ["Portafolio de Productos", ""],
        ["Categoría", "Modelos"],
        ["Livianos (urbanos/flota)", "LB (entrada de gama), LP (premium/tecnología)"],
        ["Pick-ups (agro/mixto rural)", "PB (cabina simple básica), PM (cabina doble media), PP (superior/alta motorización)"],
        ["", ""],
        ["Índice de Pestañas", ""],
        ["Pestaña", "Descripción"],
        ["Readme_Portada", "Carátula y resumen ejecutivo"],
        ["Macro_y_TipoCambio", "Supuestos macroeconómicos y matrices de Tipo de Cambio"],
        ["Capacidad_Dotacion", "Capacidades por línea/turno y dotación laboral"],
        ["Costos_Operativos", "Costos fijos y variables (ARS y USD)"],
        ["Demanda_y_Precios", "Precios y proyecciones de demanda"],
        ["Proyectos_Estrategicos", "Parámetros para consignas 2 a 6"],
        ["Estructura_Modelo_PL", "Especificación formal del modelo"],
        ["", ""],
        ["Guía de Integración con Python", ""],
        ["Uso", "Alimentar los scripts de optimización mediante pandas.read_excel()"],
        ["Arquitectura", "Sheets (BD) -> Python (Motor PL) -> Docs (Informe)"]
    ]
    for row in data1:
        ws1.append(row)
        
    # -------------------------------------------------------------
    # 2. Macro_y_TipoCambio
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Macro_y_TipoCambio")
    ws2.views.sheetView[0].showGridLines = True
    
    ws2.append(["Supuestos Macroeconómicos", "", "", "Factores Inflacionarios Anuales", ""])
    ws2.append(["Parámetro", "Valor", "", "Año", "Factor"])
    ws2.append(["Inflación de costos proyectada", 0.20, "", "Año 1", 1.0])
    ws2.append(["Tasa inventario playón (Base)", 0.25, "", "Año 2", 1.2])
    ws2.append(["Tasa inventario playón (Sensibilidad)", 0.30, "", "Año 3", 1.44])
    ws2.append(["Meses liquidados al año (con SAC)", 13, "", "Año 4", 1.728])
    ws2.append(["Prima exportación MERCOSUR", 0.05, "", "Año 5", 2.0736])
    ws2.append(["", "", "", "", ""])
    ws2.append(["Tipo de Cambio - Escenario Base", "", "", "Tipo de Cambio - Escenario Devaluación", ""])
    ws2.append(["Año", "ARS/USD", "", "Año", "ARS/USD"])
    ws2.append(["Año 1", 1500, "", "Año 1", 1500])
    ws2.append(["Año 2", 1750, "", "Año 2", 2200])
    ws2.append(["Año 3", 1950, "", "Año 3", 2650])
    ws2.append(["Año 4", 2350, "", "Año 4", 3100])
    ws2.append(["Año 5", 2800, "", "Año 5", 3550])

    # -------------------------------------------------------------
    # 3. Capacidad_Dotacion
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Capacidad_Dotacion")
    ws3.views.sheetView[0].showGridLines = True
    
    ws3.append(["Capacidad Línea A (Livianos)", "", "", "Capacidad Línea B (Pick-ups o Livianos)", ""])
    ws3.append(["Turno", "Unidades/año", "", "Turno", "Unidades/año"])
    ws3.append(["Mañana", 10000, "", "Mañana (Pick-ups)", 25000])
    ws3.append(["Tarde (75%)", 7500, "", "Mañana (Livianos)", 30000])
    ws3.append(["Máxima anual", 17500, "", "Tarde (Pick-ups 75%)", 18750])
    ws3.append(["", "", "", "Tarde (Livianos 75%)", 22500])
    ws3.append(["", "", "", "", ""])
    ws3.append(["Dotación Laboral por Sección", "", "", "Reglas Laborales por Turno", ""])
    ws3.append(["Sección", "Operarios", "", "Regla", "Valor"])
    ws3.append(["Chasis", 1, "", "Dotación Base Línea A (Mañana/Tarde)", 9])
    ws3.append(["Pintura", 1, "", "Dotación Base Línea B (Mañana/Tarde)", 11])
    ws3.append(["Motor", 2, "", "Adicionales Conjuntos Turno Tarde", 5])
    ws3.append(["Detalles Finales (Línea A)", 5, "", "", ""])
    ws3.append(["Detalles Finales (Línea B)", 7, "", "", ""])

    # -------------------------------------------------------------
    # 4. Costos_Operativos
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Costos_Operativos")
    ws4.views.sheetView[0].showGridLines = True
    
    ws4.append(["Sueldos (ARS)", "", "", "", "", ""])
    ws4.append(["Concepto", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5"])
    ws4.append(["Sueldo mensual", 2000000, 2400000, 2880000, 3456000, 4147200])
    ws4.append(["Costo anual por operario", 26000000, 31200000, 37440000, 44928000, 53913600])
    ws4.append(["", "", "", "", "", ""])
    ws4.append(["Costo Fijo de Encendido de Línea por Turno", "", "", "", "", "", ""])
    ws4.append(["Línea / Escenario", "Moneda", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5"])
    ws4.append(["Línea A (Livianos)", "ARS", 500000000, 600000000, 720000000, 864000000, 1036800000])
    ws4.append(["Línea A (Livianos)", "USD (Base)", 333333.33, 342857.14, 369230.77, 367659.57, 370285.71])
    ws4.append(["Línea A (Livianos)", "USD (Dev)", 333333.33, 272727.27, 271698.11, 278709.68, 292056.34])
    ws4.append(["Línea B (Pick-ups)", "ARS", 800000000, 960000000, 1152000000, 1382400000, 1658880000])
    ws4.append(["Línea B (Pick-ups)", "USD (Base)", 533333.33, 548571.43, 590769.23, 588255.32, 592457.14])
    ws4.append(["Línea B (Pick-ups)", "USD (Dev)", 533333.33, 436363.64, 434716.98, 445935.48, 467290.14])
    ws4.append(["", "", "", "", "", "", ""])
    ws4.append(["Costos Variables Unitarios (ARS Inflacionados)", "", "", "", "", "", ""])
    ws4.append(["Modelo", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5", "Notas"])
    ws4.append(["LB", 40000000, 48000000, 57600000, 69120000, 82944000, "Costo base livianos"])
    ws4.append(["LP", 43200000, 51840000, 62208000, 74649600, 89579520, "+8% sobre LB"])
    ws4.append(["PB", 57500000, 69000000, 82800000, 99360000, 119232000, "Costo base pick-ups"])
    ws4.append(["PM", 59225000, 71070000, 85284000, 102340800, 122808960, "+3% sobre PB"])
    ws4.append(["PP", 60950000, 73140000, 87768000, 105321600, 126385920, "+6% sobre PB"])
    ws4.append(["", "", "", "", "", "", ""])
    ws4.append(["Costos Variables Unitarios (USD Base)", "", "", "", "", "", ""])
    ws4.append(["Modelo", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5", "Notas"])
    ws4.append(["LB", 26666.67, 27428.57, 29538.46, 29412.77, 29622.86, "ARS inflacionado / TC Base"])
    ws4.append(["LP", 28800.00, 29622.86, 31901.54, 31765.79, 31992.69, ""])
    ws4.append(["PB", 38333.33, 39428.57, 42461.54, 42280.85, 42582.86, ""])
    ws4.append(["PM", 39483.33, 40611.43, 43735.38, 43549.28, 43860.34, ""])
    ws4.append(["PP", 40633.33, 41794.29, 45009.23, 44817.70, 45137.83, ""])
    ws4.append(["", "", "", "", "", "", ""])
    ws4.append(["Costos Variables Unitarios (USD Devaluación)", "", "", "", "", "", ""])
    ws4.append(["Modelo", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5", "Notas"])
    ws4.append(["LB", 26666.67, 21818.18, 21735.85, 22296.77, 23364.51, "ARS inflacionado / TC Dev"])
    ws4.append(["LP", 28800.00, 23563.64, 23474.72, 24080.52, 25233.67, ""])
    ws4.append(["PB", 38333.33, 31363.64, 31245.28, 32051.61, 33586.48, ""])
    ws4.append(["PM", 39483.33, 32304.55, 32182.64, 33013.16, 34594.07, ""])
    ws4.append(["PP", 40633.33, 33245.45, 33120.00, 33974.71, 35601.67, ""])

    # -------------------------------------------------------------
    # 5. Demanda_y_Precios
    # -------------------------------------------------------------
    ws5 = wb.create_sheet(title="Demanda_y_Precios")
    ws5.views.sheetView[0].showGridLines = True
    
    ws5.append(["Precios Unitarios (USD)", "", ""])
    ws5.append(["Modelo", "Precio de Lista (Local)", "Precio Exportación (+5%)"])
    ws5.append(["LB", 30000, "N/A"])
    ws5.append(["LP", 35000, "N/A"])
    ws5.append(["PB", 45000, 47250])
    ws5.append(["PM", 50000, 52500])
    ws5.append(["PP", 57500, 60375])
    ws5.append(["", "", "", "", "", ""])
    ws5.append(["Proyecciones de Demanda Local (Unidades)", "", "", "", "", ""])
    ws5.append(["Modelo", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5"])
    ws5.append(["LB", 2000, 2060.0, 2121.8, 2185.5, 2251.0])
    ws5.append(["LP", 4500, 4635.0, 4774.1, 4917.3, 5064.8])
    ws5.append(["PB", 12000, 12600.0, 13230.0, 13891.5, 14586.1])
    ws5.append(["PM", 4000, 4200.0, 4410.0, 4630.5, 4862.0])
    ws5.append(["PP", 6000, 6300.0, 6615.0, 6945.8, 7293.0])
    ws5.append(["", "", "", "", "", ""])
    ws5.append(["Proyecciones de Demanda Exportación MERCOSUR (Unidades)", "", "", "", "", ""])
    ws5.append(["Modelo", "Año 1", "Año 2", "Año 3", "Año 4", "Año 5"])
    ws5.append(["PB", 4000, 4000, 4000, 4000, 4000])
    ws5.append(["PM", 7500, 7500, 7500, 7500, 7500])
    ws5.append(["PP", 5000, 5000, 5000, 5000, 5000])
    ws5.append(["", "", "", "", "", ""])
    ws5.append(["Contratos y Proyectos Especiales", "", ""])
    ws5.append(["Proyecto", "Detalle", "Impacto"])
    ws5.append(["Contrato Autonomy", "1.000 LB anuales (Años 1-5)", "Precio fijo: USD 27.500"])
    ws5.append(["Propuesta Agronegocios", "1.500 PP anuales (Años 2-5)", "Precio de lista: USD 57.500"])
    ws5.append(["Boom Agropecuario", "+20% demanda Pick-ups", "Aplica Local y Exportación"])

    # -------------------------------------------------------------
    # 6. Proyectos_Estrategicos
    # -------------------------------------------------------------
    ws6 = wb.create_sheet(title="Proyectos_Estrategicos")
    ws6.views.sheetView[0].showGridLines = True
    
    ws6.append(["Consigna 2: Evaluación Devaluación", "Valor / Detalle", "", "Consigna 5: Contrato Autonomy", "Valor / Detalle"])
    ws6.append(["Escenario", "Base vs Devaluación Acelerada", "", "Objetivo", "Hallar precio umbral de conveniencia y costo de oportunidad"])
    ws6.append(["", "", "", "", ""])
    ws6.append(["Consigna 3: Reequipamiento Línea A", "Valor / Detalle", "", "Consigna 6: Costo de Playón", "Valor / Detalle"])
    ws6.append(["Inversión (Rango)", "USD 12M a 18M", "", "Tasa de sensibilidad", "30% anual"])
    ws6.append(["Parada de Planta", "1 año (Línea A cap = 0)", "", "Objetivo", "Barrido continuo de tasa de inventario"])
    ws6.append(["", "", "", "", ""])
    ws6.append(["Consigna 4: Importación China", "Valor / Detalle", "", "", ""])
    ws6.append(["Costo fijo único (Año 2)", 8000000, "", "", ""])
    ws6.append(["Costo adquisición CIF", "LB: USD 25k, LP: USD 28k", "", "", ""])
    ws6.append(["Riesgo regulatorio", "Cierre importaciones en Año 3", "", "", ""])

    # -------------------------------------------------------------
    # 7. Estructura_Modelo_PL
    # -------------------------------------------------------------
    ws7 = wb.create_sheet(title="Estructura_Modelo_PL")
    ws7.views.sheetView[0].showGridLines = True
    
    ws7.append(["Conjuntos e Índices", "Definición"])
    ws7.append(["T", "{1, 2, 3, 4, 5} (Años)"])
    ws7.append(["M", "{LB, LP, PB, PM, PP} (Modelos)"])
    ws7.append(["L", "{A, B} (Líneas)"])
    ws7.append(["K", "{Mañana, Tarde} (Turnos)"])
    ws7.append(["D", "{Local, Export, Autonomy, Agro} (Mercados)"])
    ws7.append(["", ""])
    ws7.append(["Variables de Decisión", "Tipo y Significado"])
    ws7.append(["X(m, l, k, t)", "Unidades producidas (Continua >= 0)"])
    ws7.append(["Y(l, k, t)", "1 si línea opera turno k en año t (Binaria)"])
    ws7.append(["W_tarde(t)", "1 si alguna línea opera turno tarde (Binaria)"])
    ws7.append(["S_local(m,t), S_export(m,t)", "Ventas anuales por mercado (Continua >= 0)"])
    ws7.append(["S_autonomy(t), S_agro(t)", "Ventas contratos especiales (Continua >= 0)"])
    ws7.append(["I(m, t)", "Inventario final al término del año t (Continua >= 0)"])
    ws7.append(["", ""])
    ws7.append(["Función Objetivo", "Maximizar Utilidad Neta Total = Ingresos USD - Costos USD"])
    ws7.append(["", ""])
    ws7.append(["Restricciones Principales", "Ecuación / Lógica"])
    ws7.append(["1. Balance de Inventario", "I(m, t-1) + Sum(X) = Ventas(m, t) + I(m, t)"])
    ws7.append(["2. Capacidad", "Por Línea y Turno (Ratio 1.2 en Línea B para Pick-ups)"])
    ws7.append(["3. Activación Turnos", "Vinculación Producción-Turno (X <= Cap * Y)"])
    ws7.append(["4. Límite Demanda", "Ventas <= Proyección por mercado"])
    ws7.append(["5. Contratos", "Cláusulas Autonomy y Agronegocios"])
    ws7.append(["6. Dotación Sindical", "W_tarde(t) >= Y(l, 'Tarde', t)"])

    # Auto-adjust column widths
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    val_str = str(cell.value)
                    max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
            
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    wb.save(filepath)
    print(f"Excel file successfully created at {filepath}")

if __name__ == "__main__":
    import sys
    out_path = sys.argv[1] if len(sys.argv) > 1 else "AutoITBA_Parametros.xlsx"
    build_excel(out_path)
